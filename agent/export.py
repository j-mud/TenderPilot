"""Export a proposal draft to Word, with the compliance matrix as a table."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

from .schemas import ProposalDraft, Qualification, TenderExtraction

STATUS_COLORS = {
    "met": RGBColor(0x1B, 0x7F, 0x3B),
    "partial": RGBColor(0xB2, 0x6B, 0x00),
    "gap": RGBColor(0xB0, 0x1E, 0x1E),
    "unknown": RGBColor(0x55, 0x55, 0x55),
}


def proposal_to_docx(draft: ProposalDraft, ext: TenderExtraction, q: Qualification, path: Path) -> Path:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(f"Technical proposal — {ext.title}", level=0)
    doc.add_paragraph(f"Tender: {ext.tender_id}   |   Issuer: {ext.issuer}   |   Deadline: {ext.deadline}")
    doc.add_paragraph(
        f"DRAFT r{draft.revision} prepared by TenderPilot for human review. "
        f"Recommendation: {q.decision} (fit {q.fit_score}/100)."
    ).runs[0].italic = True

    for section in draft.sections:
        doc.add_heading(section.title, level=1)
        for para in section.content.split("\n"):
            if para.strip():
                doc.add_paragraph(para.strip())

    doc.add_heading("Compliance matrix", level=1)
    checks = {c.requirement_id: c for c in q.checks}
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for cell, label in zip(table.rows[0].cells, ["ID", "Requirement", "Mandatory", "Status", "Evidence / note"]):
        cell.text = label
    for r in ext.requirements:
        c = checks.get(r.id)
        row = table.add_row().cells
        row[0].text = r.id
        row[1].text = f"{r.text} ({r.source})"
        row[2].text = "Yes" if r.mandatory else "No"
        status = c.status if c else "unknown"
        run = row[3].paragraphs[0].add_run(status.upper())
        run.bold = True
        run.font.color.rgb = STATUS_COLORS[status]
        row[4].text = f"{c.evidence} — {c.note}" if c else ""

    if q.blockers:
        doc.add_heading("Blockers to resolve before submission", level=1)
        for b in q.blockers:
            doc.add_paragraph(b, style="List Bullet")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
