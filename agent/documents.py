"""Loading tender documents and the company knowledge base.

A tender can be either:
  - a single file:   data/<set>/tenders/T101_name.pdf   (or .md/.txt/.docx/.xlsx)
  - a folder (pack): data/<set>/tenders/<tender_id>/   containing several files,
                     e.g. the main booklet, appendices and a bill of quantities.

Optional metadata copied from the Etimad listing page (deadline, value, link)
can be added as  <tender_id>.meta.json  next to a file, or  meta.json  inside a
pack folder. Values there override what the model extracts.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

DOC_EXTS = {".md", ".txt", ".pdf", ".docx", ".xlsx"}
CHUNK_CHARS = 60_000  # ~15k tokens per extraction call


@dataclass
class Tender:
    tender_id: str
    text: str
    files: list[str]
    meta: dict = field(default_factory=dict)


# ---------------------------------------------------------------- readers
def _read_pdf(path: Path) -> str:
    try:
        import pymupdf as fitz
    except ImportError:  # older PyMuPDF
        import fitz

    pages = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if len(text) < 20:  # probably a scanned page: try OCR if available
                text = _ocr_page(page) or text
            pages.append(f"[{path.name} — page {i}]\n{text}")
    return "\n\n".join(pages)


def _ocr_page(page) -> str:
    try:
        import io

        import pytesseract
        from PIL import Image

        pix = page.get_pixmap(dpi=200)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(img, lang="ara+eng")
    except Exception:
        return ""


def _read_docx(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for t_i, table in enumerate(doc.tables, start=1):
        parts.append(f"[table {t_i}]")
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _read_xlsx(path: Path, max_rows: int = 400) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"[sheet {ws.title}]")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                parts.append(f"... (truncated after {max_rows} rows)")
                break
            cells = [str(c) for c in row if c is not None]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_document(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".xlsx":
        return _read_xlsx(path)
    return path.read_text(encoding="utf-8", errors="replace")


# ---------------------------------------------------------------- tenders
def _load_meta(path: Path) -> dict:
    if path.exists():
        try:
            return json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return {}
    return {}


def load_tenders(folder: Path) -> dict[str, Tender]:
    tenders: dict[str, Tender] = {}
    if not folder.exists():
        return tenders
    for path in sorted(folder.iterdir()):
        if path.name.startswith(".") or path.name.endswith(".meta.json"):
            continue
        if path.is_dir():
            files = sorted(p for p in path.rglob("*") if p.suffix.lower() in DOC_EXTS and not p.name.startswith("."))
            if not files:
                continue
            text = "\n\n".join(f"=== FILE: {p.relative_to(path)} ===\n{read_document(p)}" for p in files)
            tenders[path.name] = Tender(path.name, text, [p.name for p in files], _load_meta(path / "meta.json"))
        elif path.suffix.lower() in DOC_EXTS:
            meta = _load_meta(path.with_name(f"{path.stem}.meta.json"))
            tenders[path.stem] = Tender(path.stem, read_document(path), [path.name], meta)
    return tenders


def chunk_text(text: str, max_chars: int | None = None) -> list[str]:
    """Split long tender text on page/file boundaries into model-sized chunks."""
    max_chars = max_chars or CHUNK_CHARS
    if len(text) <= max_chars:
        return [text]
    blocks = re.split(r"(?=\n\[[^\]\n]+ — page \d+\]|\n=== FILE: )", text)
    chunks, current = [], ""
    for block in blocks:
        pieces = [block[i : i + max_chars] for i in range(0, len(block), max_chars)] or [""]
        for piece in pieces:
            if current and len(current) + len(piece) > max_chars:
                chunks.append(current)
                current = ""
            current += piece
    if current.strip():
        chunks.append(current)
    return chunks


# ---------------------------------------------------------------- dates
_HIJRI_HINT = re.compile(r"(هـ|ه\b|AH\b|hijri)", re.I)


def normalize_deadline(value: str) -> tuple[str, str]:
    """Return (gregorian YYYY-MM-DD or '', note). Converts Hijri dates (year 13xx-15xx)."""
    if not value:
        return "", "No deadline found"
    value = value.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789"))  # Arabic-Indic digits
    m = re.search(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", value) or None
    if m:
        y, mo, d = (int(x) for x in m.groups())
    else:
        m = re.search(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", value)
        if not m:
            return "", f"Could not read deadline '{value}'"
        d, mo, y = (int(x) for x in m.groups())
    if 1300 <= y <= 1600 or (_HIJRI_HINT.search(value) and y < 1700):
        try:
            from hijridate import Hijri

            g = Hijri(y, mo, d).to_gregorian()
            return g.isoformat(), f"Converted from Hijri {y}-{mo:02d}-{d:02d}"
        except Exception as err:
            return "", f"Invalid Hijri date '{value}': {err}"
    try:
        return date(y, mo, d).isoformat(), ""
    except ValueError:
        return "", f"Invalid date '{value}'"


# ---------------------------------------------------------------- company
ANSWERS_FILE = "bid_manager_answers.md"
ANSWERS_HEADER = ("# Answers from the bid manager\n"
                  "Given after the review, so newer than the other company files.\n")
_DATE = re.compile(r"\d{4}[-/.]\d{1,2}[-/.]\d{1,2}(?:\s*هـ)?|\d{1,2}[-/.]\d{1,2}[-/.]\d{4}(?:\s*هـ)?")


class CompanyKnowledge:
    """The company's own files. Small enough to send to the model in full."""

    def __init__(self, folder: Path):
        self.folder = folder
        self.files: dict[str, str] = {
            p.name: p.read_text(encoding="utf-8")
            for p in sorted(folder.iterdir())
            if p.suffix.lower() in {".md", ".txt", ".json"}
        } if folder.exists() else {}
        certs_path = folder / "certifications.json"
        self.certs = json.loads(certs_path.read_text(encoding="utf-8")) if certs_path.exists() else {}

    def context(self) -> str:
        return "\n\n".join(f"=== FILE: {name} ===\n{text}" for name, text in self.files.items())

    def file_names(self) -> list[str]:
        return list(self.files)

    def unfinished(self) -> list[str]:
        """Template files still containing TODO placeholders."""
        return [name for name, text in self.files.items() if "TODO" in text]

    def expiring_before(self, deadline: str) -> list[dict]:
        """Certificates that will have expired by the tender deadline."""
        try:
            due = date.fromisoformat(deadline[:10])
        except ValueError:
            return []
        out = []
        for c in self.certs.get("certifications", []):
            try:
                if date.fromisoformat(c["valid_until"]) < due:
                    out.append(c)
            except (KeyError, ValueError):
                continue
        return out

    def not_held(self) -> list[str]:
        return self.certs.get("not_held", [])

    def add_answer(self, tender_id: str, question: str, answer: str) -> list[str]:
        """A bid manager's answer joins the knowledge for this run (files on disk are not changed).

        Returns the certificate updates it caused, e.g. a renewal with a new expiry date.
        """
        entry = f"\n- [{tender_id}] Question: {question}\n  Answer: {answer}\n"
        self.files[ANSWERS_FILE] = self.files.get(ANSWERS_FILE, ANSWERS_HEADER) + entry
        return self._apply_renewal(answer)

    def _apply_renewal(self, answer: str) -> list[str]:
        """'Zakat renewed, valid until 2027-12-01' -> that certificate's new expiry date.

        Only when unambiguous: exactly one certificate named (by its keyword, as a whole word)
        and exactly one date. Anything else stays a plain statement in the answers file.
        """
        dates = {d for d in (normalize_deadline(m)[0] for m in _DATE.findall(answer)) if d}
        certs = [c for c in self.certs.get("certifications", [])
                 if c.get("keyword") and re.search(rf"\b{re.escape(c['keyword'])}\b", answer, re.I)]
        if len(dates) != 1 or len(certs) != 1:
            return []
        cert, new = certs[0], dates.pop()
        old, cert["valid_until"] = cert.get("valid_until", "?"), new
        self.files["certifications.json"] = json.dumps(self.certs, ensure_ascii=False, indent=2)
        return [f"{cert['name']}: valid until {old} -> {new} (from the bid manager's answer)"]
